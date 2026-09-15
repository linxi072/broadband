# -*- coding: utf-8 -*-
"""轻量 Markdown → DOCX 转换器（中文排版：微软雅黑标题 + 宋体正文 + Table Grid 表格）
用法: python md2docx.py <input.md> <output.docx>
"""
import re
import sys

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

TITLE_COLOR = RGBColor(0x1F, 0x4E, 0x79)
GRAY = RGBColor(0x55, 0x55, 0x55)
CODE_FONT = 'Consolas'

INLINE_RE = re.compile(r'(\*\*.+?\*\*|`[^`]+`)')


def set_run(run, name, size, bold=None, color=None, ea=None):
    run.font.name = name
    run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = rpr.makeelement(qn('w:rFonts'), {})
        rpr.append(rf)
    rf.set(qn('w:ascii'), name)
    rf.set(qn('w:hAnsi'), name)
    rf.set(qn('w:eastAsia'), ea or name)


def add_rich(p, text, name, size, bold=False, color=None):
    """解析 **粗体** 与 `代码` 后写入段落"""
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**') and len(part) > 4:
            r = p.add_run(part[2:-2])
            set_run(r, name, size, True, color)
        elif part.startswith('`') and part.endswith('`') and len(part) > 2:
            r = p.add_run(part[1:-1])
            set_run(r, CODE_FONT, size - 0.5, bold, RGBColor(0xC0, 0x39, 0x2B))
        else:
            r = p.add_run(part)
            set_run(r, name, size, bold, color)


def shade(cell, hex_color):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.makeelement(qn('w:shd'), {})
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcpr.append(shd)


def convert(md_path, docx_path):
    with open(md_path, encoding='utf-8') as f:
        lines = f.read().split('\n')

    d = docx.Document()
    sec = d.sections[0]
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin = sec.bottom_margin = Cm(2.0)

    normal = d.styles['Normal']
    normal.font.name = '宋体'
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

    i = 0
    n = len(lines)
    while i < n:
        line = lines[i].rstrip()

        # 表格
        if line.startswith('|') and i + 1 < n and re.match(r'^\|[\s:|-]+\|$', lines[i + 1].strip()):
            block = []
            while i < n and lines[i].strip().startswith('|'):
                block.append(lines[i].strip())
                i += 1
            rows = []
            for idx, bl in enumerate(block):
                cells = [c.strip() for c in bl.strip('|').split('|')]
                if re.match(r'^[\s:|-]+$', ''.join(cells)):
                    continue
                rows.append(cells)
            if rows:
                ncol = max(len(r) for r in rows)
                t = d.add_table(rows=len(rows), cols=ncol)
                t.style = 'Table Grid'
                t.alignment = WD_TABLE_ALIGNMENT.CENTER
                for ri, row in enumerate(rows):
                    for ci in range(ncol):
                        cell = t.cell(ri, ci)
                        txt = row[ci] if ci < len(row) else ''
                        txt = txt.replace('<br>', '\n')
                        cell.text = ''
                        p = cell.paragraphs[0]
                        if ri == 0:
                            shade(cell, '1F4E79')
                            add_rich(p, txt.strip('*'), '微软雅黑', 9.5, True, RGBColor(0xFF, 0xFF, 0xFF))
                        else:
                            if ri % 2 == 0:
                                shade(cell, 'F2F6FA')
                            add_rich(p, txt, '宋体', 9.5)
                d.add_paragraph()
            continue

        if not line.strip():
            i += 1
            continue

        if line.startswith('# '):
            p = d.add_paragraph()
            p.style = d.styles['Title']
            add_rich(p, line[2:].strip(), '微软雅黑', 20, True, TITLE_COLOR)
            i += 1
            continue
        if line.startswith('## '):
            p = d.add_paragraph()
            p.style = d.styles['Heading 1']
            add_rich(p, line[3:].strip(), '微软雅黑', 15, True, TITLE_COLOR)
            i += 1
            continue
        if line.startswith('### '):
            p = d.add_paragraph()
            p.style = d.styles['Heading 2']
            add_rich(p, line[4:].strip(), '微软雅黑', 12.5, True, TITLE_COLOR)
            i += 1
            continue
        if re.match(r'^(-{3,}|\*{3,})$', line.strip()):
            i += 1
            continue

        # 引用块（合并连续 > 行）
        if line.startswith('>'):
            buf = []
            while i < n and lines[i].startswith('>'):
                buf.append(re.sub(r'^>\s?', '', lines[i]).rstrip())
                i += 1
            text = '\n'.join(x for x in buf if x.strip())
            p = d.add_paragraph()
            add_rich(p, text, '宋体', 9.5, False, GRAY)
            p.paragraph_format.left_indent = Cm(0.4)
            continue

        # 列表
        m = re.match(r'^(\s*)[-*]\s+(.*)$', line)
        if m:
            p = d.add_paragraph(style='List Bullet')
            add_rich(p, m.group(2), '宋体', 10.5)
            i += 1
            continue
        m = re.match(r'^(\s*)(\d+)\.\s+(.*)$', line)
        if m:
            p = d.add_paragraph(style='List Number')
            add_rich(p, m.group(3), '宋体', 10.5)
            i += 1
            continue

        p = d.add_paragraph()
        add_rich(p, line, '宋体', 10.5)
        i += 1

    d.save(docx_path)
    print('saved', docx_path)


if __name__ == '__main__':
    convert(sys.argv[1], sys.argv[2])
